"""Excel export structure and a smoke test of every app page (streamlit AppTest).

The numeric check of the Excel formulas needs Excel itself:
    python scripts/export_workbook.py --verify
"""

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from scripts.export_workbook import LSL, USL, load
from stackup import controlplan, stack
from stackup.excel import workbook_bytes

APP = str(Path(__file__).resolve().parents[1] / "app.py")


@pytest.fixture(scope="module")
def wb():
    dims, modes, plan = load()
    mc = stack.monte_carlo(dims, LSL, USL, n=1000)
    findings = controlplan.check_stack_assumptions(dims, plan)
    return load_workbook(BytesIO(workbook_bytes(dims, LSL, USL, mc, 1000, modes, plan, findings)))


def test_sheets(wb):
    assert wb.sheetnames == ["Stack-up", "Monte Carlo", "PFMEA", "Control Plan", "Checks"]


def test_stack_values_are_formulas(wb):
    ws = wb["Stack-up"]
    assert ws["I7"].value == "=D7+(E7+F7)/2"
    assert ws["J7"].value == "=(E7-F7)/2"
    formulas = [ws.cell(r, 2).value for r in range(12, 25)]
    assert all(isinstance(f, str) and f.startswith("=") for f in formulas)
    assert any("SQRT(SUM(M7:M10))" in f for f in formulas)
    assert any("_xlfn.NORM.DIST" in f for f in formulas)


def test_pfmea_rpn_and_priority_are_formulas(wb):
    ws = wb["PFMEA"]
    assert ws["L5"].value == "=F5*H5*K5"
    assert ws["M5"].value.startswith('=IF(F5>=9,"High"')


def test_checks_sheet_reports_ok(wb):
    assert wb["Checks"]["A4"].value == "OK"


def run(page):
    at = AppTest.from_file(APP, default_timeout=90)
    at.run()
    at.sidebar.radio(key="page").set_value(page).run()
    assert not at.exception, at.exception
    return at


def test_stack_page():
    at = run("📏 Stack-up")
    labels = {m.label: m for m in at.metric}
    assert labels["Worst case"].delta == "FAIL" and labels["RSS"].delta == "PASS"


def test_pfmea_page():
    at = run("⚠️ PFMEA")
    assert {m.label: m.value for m in at.metric}["High priority"] == "2"


def test_control_plan_page():
    run("📋 Control plan")


def test_traceability_page():
    at = run("🔗 Traceability & export")
    assert at.success and not at.error
