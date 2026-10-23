"""Export the stack-up, PFMEA and control plan to one Excel workbook with live formulas.

Engineers review these documents in Excel, so the workbook does not paste
values from Python: means, half tolerances, sigmas, worst case, RSS, Cpk,
PPM, RPN and priority are all Excel formulas.  Change a tolerance or a
rating in Excel and everything downstream recalculates.  Only the Monte
Carlo summary and the traceability findings are static (they come from Python).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .controlplan import ControlItem, Finding
from .pfmea import FailureMode
from .stack import Dimension, MonteCarloResult

ORANGE = "FF8A3D"
DARK = "14161B"
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=DARK)
INPUT_FILL = PatternFill("solid", fgColor="FFF4EC")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, labels: list[str], widths: list[int] | None = None):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row, c, label)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center"), THIN
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _stack_sheet(wb, dims: list[Dimension], lsl, usl, title: str):
    ws = wb.active
    ws.title = "Stack-up"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=ORANGE)
    ws["A2"] = "Orange cells are inputs. Everything else is a formula: edit a tolerance and the results update."
    ws["A3"], ws["B3"] = "Gap LSL", lsl
    ws["A4"], ws["B4"] = "Gap USL", usl
    for ref in ("B3", "B4"):
        ws[ref].fill = INPUT_FILL

    top = 6
    labels = ["Dimension", "Source", "Direction (+1/-1)", "Nominal", "Upper dev", "Lower dev", "Assumed Cp",
              "Distribution", "Mean", "Half tol (t)", "Sigma", "Dir x Mean", "t²", "σ²", "Worst-case %", "Statistical %"]
    _header(ws, top, labels, [28, 34, 11, 10, 10, 10, 10, 12, 10, 10, 10, 11, 10, 11, 12, 12])
    first, last = top + 1, top + len(dims)
    for i, d in enumerate(dims):
        r = first + i
        values = [d.name, d.source, d.direction, d.nominal, d.upper, d.lower, d.cp, d.distribution]
        for c, v in enumerate(values, 1):
            ws.cell(r, c, v).fill = INPUT_FILL
        ws.cell(r, 9, f"=D{r}+(E{r}+F{r})/2")
        ws.cell(r, 10, f"=(E{r}-F{r})/2")
        ws.cell(r, 11, f'=IF(H{r}="uniform",J{r}/SQRT(3),J{r}/(3*G{r}))')
        ws.cell(r, 12, f"=C{r}*I{r}")
        ws.cell(r, 13, f"=J{r}^2")
        ws.cell(r, 14, f"=K{r}^2")
        ws.cell(r, 15, f"=J{r}/SUM($J${first}:$J${last})")
        ws.cell(r, 16, f"=N{r}/SUM($N${first}:$N${last})")
        for c in range(1, 17):
            ws.cell(r, c).border = THIN
        for c in range(4, 15):
            ws.cell(r, c).number_format = "0.0000"
        for c in (15, 16):
            ws.cell(r, c).number_format = "0.0%"

    rng = lambda col: f"{col}{first}:{col}{last}"
    s = last + 2
    rows = [
        ("Nominal gap", f"=SUMPRODUCT({rng('C')},{rng('D')})", "0.0000"),
        ("Mean gap", f"=SUM({rng('L')})", "0.0000"),
        ("Worst-case ± (Σt)", f"=SUM({rng('J')})", "0.0000"),
        ("Worst-case min", f"=B{s + 1}-B{s + 2}", "0.0000"),
        ("Worst-case max", f"=B{s + 1}+B{s + 2}", "0.0000"),
        ("Worst-case verdict", f'=IF(AND(B{s + 3}>=$B$3,B{s + 4}<=$B$4),"PASS","FAIL")', None),
        ("RSS ± (√Σt²)", f"=SQRT(SUM({rng('M')}))", "0.0000"),
        ("RSS min", f"=B{s + 1}-B{s + 6}", "0.0000"),
        ("RSS max", f"=B{s + 1}+B{s + 6}", "0.0000"),
        ("RSS verdict", f'=IF(AND(B{s + 7}>=$B$3,B{s + 8}<=$B$4),"PASS","FAIL")', None),
        ("Gap sigma (√Σσ²)", f"=SQRT(SUM({rng('N')}))", "0.00000"),
        ("Predicted gap Cpk", f"=MIN($B$4-B{s + 1},B{s + 1}-$B$3)/(3*B{s + 10})", "0.00"),
        # upper tail by symmetry (P(x > USL) = P(x < 2*mean - USL)) so it isn't lost in 1 - 0.99999...
        # _xlfn. tells Excel this is the 2010+ NORM.DIST, otherwise it shows #NAME?
        ("Predicted PPM out of spec", f"=(_xlfn.NORM.DIST($B$3,B{s + 1},B{s + 10},TRUE)"
                                      f"+_xlfn.NORM.DIST(2*B{s + 1}-$B$4,B{s + 1},B{s + 10},TRUE))*1000000", "0.000"),
    ]
    for i, (label, formula, fmt) in enumerate(rows):
        r = s + i
        ws.cell(r, 1, label).font = Font(bold=True)
        cell = ws.cell(r, 2, formula)
        cell.border = THIN
        if fmt:
            cell.number_format = fmt
    for verdict_row in (s + 5, s + 9):
        ref = f"B{verdict_row}"
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN))
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED))
    ws.freeze_panes = ws.cell(first, 2)
    return {"first": first, "last": last, "summary": s}


def _mc_sheet(wb, mc: MonteCarloResult, n: int):
    ws = wb.create_sheet("Monte Carlo")
    ws["A1"] = "Monte Carlo summary (static values from Python, not live formulas)"
    ws["A1"].font = Font(bold=True, size=12, color=ORANGE)
    rows = [("Samples", n), ("Mean gap", mc.mean), ("Std dev", mc.std), ("0.135th percentile", mc.p_low),
            ("99.865th percentile", mc.p_high), ("PPM below LSL", mc.ppm_below), ("PPM above USL", mc.ppm_above)]
    for i, (label, value) in enumerate(rows, 3):
        ws.cell(i, 1, label).font = Font(bold=True)
        ws.cell(i, 2, value).number_format = "0" if label == "Samples" else "0.00000"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14


def _pfmea_sheet(wb, modes: list[FailureMode], rpn_threshold: int):
    ws = wb.create_sheet("PFMEA")
    ws["A1"] = "Process FMEA (AIAG 4th ed. scales)"
    ws["A1"].font = Font(bold=True, size=12, color=ORANGE)
    ws["A2"], ws["B2"] = "RPN action threshold", rpn_threshold
    ws["B2"].fill = INPUT_FILL
    ws["D2"] = "Priority: High if S ≥ 9; Medium if RPN ≥ threshold or (S ≥ 7 and O ≥ 4); else Low."
    top = 4
    labels = ["ID", "Process step", "Function / requirement", "Failure mode", "Effect", "S", "Cause", "O",
              "Prevention control", "Detection control", "D", "RPN", "Priority", "Class", "Recommended action"]
    _header(ws, top, labels, [6, 20, 24, 26, 32, 5, 26, 5, 26, 24, 5, 7, 10, 7, 40])
    for i, m in enumerate(modes):
        r = top + 1 + i
        values = [m.id, m.step, m.function, m.failure_mode, m.effect, m.severity, m.cause, m.occurrence,
                  m.prevention, m.detection_control, m.detection]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            if c in (6, 8, 11):
                cell.fill = INPUT_FILL
        ws.cell(r, 12, f"=F{r}*H{r}*K{r}")
        ws.cell(r, 13, f'=IF(F{r}>=9,"High",IF(OR(L{r}>=$B$2,AND(F{r}>=7,H{r}>=4)),"Medium","Low"))')
        ws.cell(r, 14, m.special_char)
        ws.cell(r, 15, m.action)
        for c in range(1, 16):
            ws.cell(r, c).border, ws.cell(r, c).alignment = THIN, WRAP
    last = top + len(modes)
    area = f"M{top + 1}:M{last}"
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"High"'], fill=RED))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"Medium"'], fill=AMBER))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"Low"'], fill=GREEN))
    ws.conditional_formatting.add(f"F{top + 1}:F{last}", FormulaRule(formula=[f"F{top + 1}>=9"], fill=RED))
    ws.freeze_panes = ws.cell(top + 1, 2)
    return {"first": top + 1, "last": last}


def _plan_sheet(wb, plan: list[ControlItem]):
    ws = wb.create_sheet("Control Plan")
    ws["A1"] = "Control plan (production)"
    ws["A1"].font = Font(bold=True, size=12, color=ORANGE)
    labels = ["Process step", "Characteristic", "Class", "LSL", "USL", "Measurement method", "Sample size",
              "Frequency", "Control method", "Required Cpk", "Reaction plan", "PFMEA ref", "Stack-up dimension"]
    _header(ws, 3, labels, [20, 30, 7, 9, 9, 26, 10, 14, 22, 10, 40, 10, 24])
    for i, it in enumerate(plan):
        r = 4 + i
        values = [it.step, it.characteristic, it.char_class, it.lsl, it.usl, it.method, it.sample_size,
                  it.frequency, it.control_method, it.required_cpk, it.reaction_plan, it.pfmea_id, it.stack_dimension]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.border, cell.alignment = THIN, WRAP
    ws.freeze_panes = "B4"


def _checks_sheet(wb, findings: list[Finding]):
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Traceability checks: PFMEA ↔ control plan ↔ stack-up assumptions"
    ws["A1"].font = Font(bold=True, size=12, color=ORANGE)
    _header(ws, 3, ["Level", "Finding"], [10, 110])
    if not findings:
        ws["A4"], ws["B4"] = "OK", "Every special characteristic and Medium/High PFMEA line is controlled, and every Cp the stack-up assumes is required by the control plan."
        ws["A4"].fill = GREEN
    for i, f in enumerate(findings):
        ws.cell(4 + i, 1, f.level.upper()).fill = RED if f.level == "error" else AMBER
        ws.cell(4 + i, 2, f.message)


def build_workbook(dims, lsl, usl, mc, mc_n, modes, plan, findings, rpn_threshold=100,
                   title="Shaft end play stack-up") -> Workbook:
    wb = Workbook()
    _stack_sheet(wb, dims, lsl, usl, title)
    _mc_sheet(wb, mc, mc_n)
    _pfmea_sheet(wb, modes, rpn_threshold)
    _plan_sheet(wb, plan)
    _checks_sheet(wb, findings)
    return wb


def workbook_bytes(*args, **kwargs) -> bytes:
    buf = BytesIO()
    build_workbook(*args, **kwargs).save(buf)
    return buf.getvalue()
